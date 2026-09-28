import cocotb
from cocotb.triggers import Timer, ReadOnly
import random
from openpyxl import load_workbook, Workbook
import time 
from datetime import datetime, timezone
from zoneinfo import ZoneInfo


@cocotb.test()
async def fullAdder(dut):
    #timestamp for file's name
    ts = time.time()
    dtUTC = datetime.fromtimestamp(ts, tz=timezone.utc)

    #excel
    wb = Workbook()
    ws = wb.active
    ws.append(["index", "A", "B", "S", "Cout", "isTrue"])
    
    wb.active.title = "sub"
    for timeVar in range(10000,1, -100):
        for i in range(0,30):
            dut.A.value = random.randint(0,2147483647)
            dut.B.value = random.randint(0,2147483647)
            dut.sub.value = 1
            await Timer(timeVar, unit="ps")
            #code U2 conversion
            if dut.carryOut.value == 0:
                x = int(dut.S.value)
                x = x - (2**32)
                print("The value of subbing : A = %d & B = %d with sub = %d is equal %d, carryOut: %d" % (dut.A.value, ~(dut.bIn.value), dut.sub.value, x, dut.carryOut.value))
                ws.append([i+1, int(dut.A.value), int(dut.B.value), x, int(dut.carryOut.value),("=IF(OR((B%d-C%d)=D%d,E%d),1,0)" % ((i+2),(i+2),(i+2),(i+2)))])
            else:
                print("The value of subbing : A = %d & B = %d with sub = %d is equal %d, carryOut: %d" % (dut.A.value, dut.bIn.value, dut.sub.value, dut.S.value, dut.carryOut.value))
                ws.append([i+1, int(dut.A.value), int(dut.B.value), int(dut.S.value), int(dut.carryOut.value),("=IF(OR((B%d-C%d)=D%d,E%d),1,0)" % ((i+2),(i+2),(i+2),(i+2)))])

    ws2 = wb.create_sheet("add")
    ws2.append(["index", "A", "B", "S", "Cout", "isTrue"])
    for timeVar in range(10000,1, -100):            
        for i in range(0,30):
            dut.A.value = random.randint(0,4294967295)
            dut.B.value = random.randint(0,4294967295)
            dut.sub.value = 0
            await Timer(timeVar, unit="ps")
            print("The value of adding : A = %d & B = %d with sub = %d is equal %d, carryOut: %d" % (dut.A.value, dut.bIn.value, dut.sub.value, dut.S.value, dut.carryOut.value))
            ws2.append([i+1, int(dut.A.value), int(dut.B.value), int(dut.S.value), int(dut.carryOut.value),("=IF(OR((B%d+C%d)=D%d,E%d),1,0)" % ((i+2),(i+2),(i+2),(i+2)))])
    
    wb.save("./Benchmarks/addSub/"+str(dtUTC)+".xlsx")
